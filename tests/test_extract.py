from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from docx import Document as DocxDocument

from documind.extract import extract_text


def test_extract_txt_returns_decoded_text():
    assert extract_text("notes.txt", b"hello world") == "hello world"


def test_extract_txt_ignores_bad_bytes():
    data = b"hello \xff\xfe world"
    text = extract_text("notes.txt", data)
    assert "hello" in text and "world" in text


def test_extract_txt_case_insensitive_extension():
    assert extract_text("NOTES.TXT", b"hi") == "hi"


def test_extract_unsupported_extension_raises():
    with pytest.raises(ValueError):
        extract_text("file.xyz", b"whatever")


def _make_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Para one.")
    doc.add_paragraph("")
    doc.add_paragraph("Para two.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "A"
    table.rows[0].cells[1].text = "B"
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_docx_includes_paragraphs_and_table():
    text = extract_text("report.docx", _make_docx_bytes())
    assert "Para one." in text
    assert "Para two." in text
    assert "A | B" in text


def test_extract_docx_case_insensitive_extension():
    text = extract_text("Report.DOCX", _make_docx_bytes())
    assert "Para one." in text


def test_extract_pdf_joins_page_text():
    page1, page2 = MagicMock(), MagicMock()
    page1.extract_text.return_value = "Page one text."
    page2.extract_text.return_value = "Page two text."
    fake_reader = MagicMock(pages=[page1, page2])
    with patch("documind.extract.PdfReader", return_value=fake_reader):
        text = extract_text("doc.pdf", b"irrelevant-bytes")
    assert text == "Page one text.\nPage two text."


def test_extract_pdf_skips_blank_pages():
    blank, real = MagicMock(), MagicMock()
    blank.extract_text.return_value = ""
    real.extract_text.return_value = "Real content."
    fake_reader = MagicMock(pages=[blank, real])
    with patch("documind.extract.PdfReader", return_value=fake_reader):
        text = extract_text("doc.pdf", b"irrelevant-bytes")
    assert text == "Real content."
