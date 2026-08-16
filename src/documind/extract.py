from io import BytesIO
from docx import Document as DocxDocument
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded .docx, .pdf, or .txt file."""
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _extract_docx(data)
    if lower.endswith(".pdf"):
        return _extract_pdf(data)
    if lower.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError(f"unsupported file type: {filename}")


def _extract_docx(data: bytes) -> str:
    doc = DocxDocument(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Include table cell text too — often where structured facts live.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(BytesIO(data))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(p for p in parts if p.strip())
